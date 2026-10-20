"""Scoped Word credits are document properties, never visible question content."""

from io import BytesIO
import zipfile

from docx import Document
from lxml import etree
from PIL import Image

from mathbank import word_export_helper as word
from mathbank.content_sources import content_sources_notice


CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"


def parts(data):
    with zipfile.ZipFile(BytesIO(data)) as package:
        assert package.testzip() is None
        assert len(package.namelist()) == len(set(package.namelist()))
        return {name: package.read(name) for name in package.namelist()}


def sample_docx():
    document = Document()
    document.add_paragraph("Visible question and ordinary braces {a_n}.")
    document.core_properties.title = "Original title"
    document.core_properties.author = "Original teacher"
    document.core_properties.subject = "Original subject"
    image = BytesIO()
    Image.new("RGB", (12, 9), "blue").save(image, format="PNG")
    image.seek(0)
    document.add_picture(image)
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def owned_property(package):
    root = etree.fromstring(package["docProps/custom.xml"])
    owned = [node for node in root if node.get("name") == "MathBankContentSources"]
    assert len(owned) == 1
    assert len(owned[0]) == 1
    assert owned[0][0].tag == f"{{{VT}}}lpwstr"
    return root, owned[0]


def test_docx_property_preserves_every_body_style_media_and_original_property_byte():
    source = sample_docx()
    notice = content_sources_notice()
    result = word._with_content_sources_docx_property(source, notice)
    before, after = parts(source), parts(result)
    allowed = {"[Content_Types].xml", "_rels/.rels", "docProps/custom.xml"}
    assert all(after[name] == value for name, value in before.items() if name not in allowed)
    assert set(after) - set(before) == {"docProps/custom.xml"}
    _, prop = owned_property(after)
    assert prop[0].text == notice
    assert prop.get("fmtid") == FMTID
    assert notice not in after["word/document.xml"].decode("utf-8")
    assert [p.text for p in Document(BytesIO(result)).paragraphs] == [p.text for p in Document(BytesIO(source)).paragraphs]
    relationships = etree.fromstring(after["_rels/.rels"])
    assert sum(node.get("Type") == REL_TYPE for node in relationships) == 1
    types = etree.fromstring(after["[Content_Types].xml"])
    assert sum(node.get("PartName") == "/docProps/custom.xml" and node.get("ContentType") == CONTENT_TYPE for node in types) == 1
    assert word._with_content_sources_docx_property(result, notice) == result


def existing_custom_docx():
    values = parts(sample_docx())
    root = etree.Element(f"{{{CUSTOM}}}Properties", nsmap={None: CUSTOM, "vt": VT})
    integer = etree.SubElement(root, f"{{{CUSTOM}}}property", fmtid=FMTID, pid="7", name="ExistingInteger")
    etree.SubElement(integer, f"{{{VT}}}i4").text = "42"
    text = etree.SubElement(root, f"{{{CUSTOM}}}property", fmtid=FMTID, pid="2", name="ExistingTeacher")
    etree.SubElement(text, f"{{{VT}}}lpwstr").text = "原教师 & 原说明"
    values["docProps/custom.xml"] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
    relationships = etree.fromstring(values["_rels/.rels"])
    etree.SubElement(relationships, f"{{{RELS}}}Relationship", Id="rId99", Type=REL_TYPE, Target="docProps/custom.xml")
    values["_rels/.rels"] = etree.tostring(relationships, encoding="UTF-8", xml_declaration=True)
    types = etree.fromstring(values["[Content_Types].xml"])
    etree.SubElement(types, f"{{{TYPES}}}Override", PartName="/docProps/custom.xml", ContentType=CONTENT_TYPE)
    values["[Content_Types].xml"] = etree.tostring(types, encoding="UTF-8", xml_declaration=True)
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as package:
        for name, value in values.items():
            package.writestr(name, value)
    return result.getvalue()


def test_existing_custom_properties_relationships_and_content_types_are_preserved():
    source = existing_custom_docx()
    before = parts(source)
    result = word._with_content_sources_docx_property(source, "限定来源\nhttps://example.invalid/\n")
    after = parts(result)
    assert all(after[name] == value for name, value in before.items() if name != "docProps/custom.xml")
    before_root = etree.fromstring(before["docProps/custom.xml"])
    root, own = owned_property(after)
    for previous in before_root:
        current = next(node for node in root if node.get("name") == previous.get("name"))
        assert etree.tostring(current) == etree.tostring(previous)
    assert own.get("pid") not in {"2", "7"}
    assert len({node.get("pid") for node in root}) == len(root)
    reopened = Document(BytesIO(result))
    assert reopened.core_properties.author == "Original teacher"
    assert word._with_content_sources_docx_property(result, own[0].text) == result
    changed = word._with_content_sources_docx_property(result, "Updated scoped source notice")
    _, updated = owned_property(parts(changed))
    assert updated.get("pid") == own.get("pid")
    assert updated[0].text == "Updated scoped source notice"


def test_real_build_final_bytes_cover_single_docx_and_both_bundle_documents(monkeypatch):
    raw = []
    save = word.WordExamBuilder.save_bytes
    def capture(builder):
        value = save(builder)
        raw.append(value)
        return value
    monkeypatch.setattr(word.WordExamBuilder, "save_bytes", capture)
    output, diagnostics = word.build_word_document("来源测试", "", "exam", [
        {"question": {"id": 1, "content": "Visible original question.", "question_type": "fill_in_blank"}, "score": 5}
    ], show_secret=False, show_notice=False)
    assert len(raw) == 1
    before, after = parts(raw[0]), parts(output)
    for name in before:
        if name not in {"[Content_Types].xml", "_rels/.rels"}:
            assert before[name] == after[name]
    assert owned_property(after)[1][0].text == content_sources_notice()
    assert not diagnostics["failed_formulas"]
    assert "Visible original question." in " ".join(p.text for p in Document(BytesIO(output)).paragraphs)
    bundled = word.create_word_bundle_zip("来源测试", output, output)
    with zipfile.ZipFile(BytesIO(bundled)) as package:
        for name in ["来源测试.docx", "来源测试_含答案与解析.docx"]:
            assert owned_property(parts(package.read(name)))[1][0].text == content_sources_notice()
        assert package.namelist().count("CONTENT_SOURCES.txt") == 1
        assert package.namelist().count("content_sources.json") == 1
