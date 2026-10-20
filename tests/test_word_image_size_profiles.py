"""Actual DOCX geometry for explicit body, table and choice image profiles."""
from io import BytesIO

from docx import Document
from docx.shared import Cm
from PIL import Image
import pytest

from mathbank.image_layout import (
    BODY_IMAGE_SIZE_CM, COMPACT_CHOICE_IMAGE_SIZE_CM, TABLE_IMAGE_MAX_HEIGHT_CM,
)
from mathbank.word_export_helper import WordExamBuilder, WordExportDiagnostics


def _builder():
    return WordExamBuilder({}, WordExportDiagnostics())


def _image(tmp_path, name, dimensions=(200, 1200)):
    path = tmp_path / name
    Image.new("RGB", dimensions, "white").save(path)
    return path


def _written_shapes(builder):
    """Save/reopen the real package so assertions cover persisted DrawingML."""
    stream = BytesIO()
    builder.doc.save(stream)
    stream.seek(0)
    return list(Document(stream).inline_shapes)


def _geometry_cm(shape):
    return (int(shape.width) / int(Cm(1)), int(shape.height) / int(Cm(1)))


def test_explicit_large_body_image_keeps_full_height_and_aspect_ratio(tmp_path):
    image = _image(tmp_path, "large-body.png")
    builder = _builder()
    builder.add_mixed(
        builder.doc.add_paragraph(), "前文 ![](/static/uploads/large-body.png) 后文",
        images=[image], image_layouts={image.name: {"align": "left", "size": "large"}},
    )
    shapes = _written_shapes(builder)
    assert len(shapes) == 1
    width, height = _geometry_cm(shapes[0])
    assert height == pytest.approx(BODY_IMAGE_SIZE_CM["large"][1], abs=0.001)
    assert height > 2.8 * 2.54  # the former compatibility cap truncated large
    assert width / height == pytest.approx(1 / 6, abs=0.001)
    assert builder.diagnostics.missing_images == 0


def test_explicit_cell_height_still_caps_large_body_image(tmp_path):
    image = _image(tmp_path, "bounded-body.png")
    builder = _builder()
    builder.add_mixed(
        builder.doc.add_paragraph(), "![](/static/uploads/bounded-body.png)",
        images=[image], image_layouts={image.name: {"align": "center", "size": "large"}},
        max_image_height=1.0,
    )
    assert _geometry_cm(_written_shapes(builder)[0])[1] == pytest.approx(2.54, abs=0.001)


@pytest.mark.parametrize("table_type", ["markdown", "tabular"])
@pytest.mark.parametrize("size", ["auto", "large"])
def test_table_images_respect_four_cm_height_for_all_layout_sizes(tmp_path, table_type, size):
    image = _image(tmp_path, "table-tall.png")
    builder = _builder()
    layouts = {image.name: {"align": "right", "size": size}}
    reference = "![](/static/uploads/table-tall.png)"
    if table_type == "markdown":
        builder.add_markdown_table("| 图 | 说明 |\n|---|---|\n| " + reference + " | 图形 |", images=[image], image_layouts=layouts)
    else:
        builder.add_tabular("c", reference, images=[image], image_layouts=layouts)
    shapes = _written_shapes(builder)
    assert len(shapes) == 1
    width, height = _geometry_cm(shapes[0])
    assert height == pytest.approx(TABLE_IMAGE_MAX_HEIGHT_CM, abs=0.001)
    assert width / height == pytest.approx(1 / 6, abs=0.001)
    assert builder.diagnostics.missing_images == 0


@pytest.mark.parametrize("dimensions", [(1200, 200), (200, 1200), (600, 600)])
def test_four_pure_image_choices_use_three_cm_square_caps(tmp_path, dimensions):
    images = [_image(tmp_path, f"option-{label}.png", dimensions) for label in "ABCD"]
    builder = _builder()
    choices = [f"![](/static/uploads/{image.name})" for image in images]
    builder.add_choices(choices, images=images)
    shapes = _written_shapes(builder)
    assert len(shapes) == 4
    for shape in shapes:
        width, height = _geometry_cm(shape)
        assert width <= COMPACT_CHOICE_IMAGE_SIZE_CM[0] + 0.001
        assert height <= COMPACT_CHOICE_IMAGE_SIZE_CM[1] + 0.001
        assert max(width, height) == pytest.approx(3.0, abs=0.001)
        assert width / height == pytest.approx(dimensions[0] / dimensions[1], rel=0.001)
    assert builder.diagnostics.missing_images == 0


def test_authored_choice_layout_is_preserved_instead_of_compact_override(tmp_path):
    images = [_image(tmp_path, f"authored-{label}.png") for label in "ABCD"]
    builder = _builder()
    choices = [f"![](/static/uploads/{image.name})" for image in images]
    builder.add_choices(
        choices, images=images,
        image_layouts={images[0].name: {"align": "right", "size": "small"}},
    )
    shapes = _written_shapes(builder)
    assert len(shapes) == 4
    assert _geometry_cm(shapes[0])[1] == pytest.approx(BODY_IMAGE_SIZE_CM["small"][1], abs=0.001)
    assert _geometry_cm(shapes[0])[1] > COMPACT_CHOICE_IMAGE_SIZE_CM[1]
    drawing_paragraphs = builder.doc.element.xpath(".//w:p[.//w:drawing]")
    assert drawing_paragraphs[0].xpath("./w:pPr/w:jc/@w:val") == ["right"]
    assert builder.diagnostics.missing_images == 0
