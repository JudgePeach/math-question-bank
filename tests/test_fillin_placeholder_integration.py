"""Answer-slot decoration is canonicalized across source and export copies."""

import ast
from io import BytesIO
import re

from docx import Document

from mathbank.content_locks import lock_visible_math, reconcile_visible_math
from mathbank.docx_helper import extract_docx_markdown
from mathbank.fillin_placeholders import normalize_fillin_placeholder_clusters
from mathbank.paper_helper import clean_content_for_latex
from mathbank.paths import PROJECT_ROOT
from mathbank.source_metadata import normalize_source_fillin
from mathbank.word_export_helper import BlankToken, tokenize_mixed_content


CLUSTER = r"\fillin\underline{▲}\fillin"


def _normalizer():
    tree = ast.parse((PROJECT_ROOT / "main.py").read_text())
    function = next(node for node in tree.body if isinstance(node, ast.FunctionDef)
                    and node.name == "normalize_fillin_macro")
    namespace = {"re": re}
    exec(compile(ast.Module(body=[function], type_ignores=[]), "main.py", "exec"), namespace)
    return namespace["normalize_fillin_macro"]


def test_ocr_and_storage_normalizer_merge_old_blank_glyphs_after_their_conversion():
    normalize = _normalizer()
    for slot in (CLUSTER, r"____\underline{▲}____", r"\_\_\_\_\underline{▲}\_\_\_\_",
                 r"$\fillin$\underline{▲}$\fillin$"):
        assert normalize("结果为" + slot + "．") == r"结果为\fillin．"
    protected = r"几何▲及$\underline{AB}$，`" + CLUSTER + "`。"
    assert normalize(protected) == protected


def test_native_word_runs_are_repaired_after_source_preservation_without_changing_extraction(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("1. 结果为")
    for value in ("    ", "▲", "    "):
        paragraph.add_run(value).underline = True
    paragraph.add_run("．")
    output = BytesIO()
    document.save(output)
    extracted = extract_docx_markdown(output.getvalue(), output_dir=tmp_path)
    assert extracted["success"] is True
    raw = extracted["markdown"]
    assert CLUSTER in raw
    expected = raw.replace(CLUSTER, r"\fillin")
    assert normalize_source_fillin(raw, _normalizer()) == expected
    assert normalize_source_fillin(expected, _normalizer()) == expected
    assert extracted["markdown"] == raw


def test_source_matching_treats_only_complete_triangle_slots_as_layout():
    source = "1. 已知 $x=1$，结果为" + CLUSTER + "．"
    _, locks = lock_visible_math(source, "triangle-slot")
    questions = [{"content": r"已知 $x=1$，结果为\fillin．", "answer_markdown": ""}]
    report = reconcile_visible_math(questions, locks, source)
    assert report["math_locks_missing"] == report["source_review_count"] == 0
    assert not report["unmatched_source"]
    assert source.endswith(CLUSTER + "．")


def test_pdf_and_word_exports_have_one_answer_line_and_keep_real_underlines():
    source = "结果为" + CLUSTER + r"．另有$\underline{a_n}$。"
    assert clean_content_for_latex(source) == r"结果为\fillin．另有$\underline{a_n}$。"
    tokens = tokenize_mixed_content(source)
    assert sum(isinstance(token, BlankToken) for token in tokens) == 1
    assert "▲" not in "".join(token for token in tokens if isinstance(token, str))
    assert CLUSTER in source


def test_non_placeholder_triangle_and_underline_remain_exact():
    for value in (r"甲为▲，乙为△。", r"\underline{▲}", r"\underline{AB}",
                  r"\fillin[2cm]\underline{▲}\fillin", r"$\fillin\underline{▲}\fillin$"):
        assert normalize_fillin_placeholder_clusters(value) == value
