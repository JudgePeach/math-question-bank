"""Finite, physical answer space can continue across paper pages."""
from io import BytesIO
import os
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING
import pytest

from mathbank.paper_layout import item_solution_space_cm, normalize_solution_space_cm
from mathbank.paper_helper import build_latex_document, compile_tex_to_pdf
from mathbank.word_export_helper import build_word_document
from test_paper_pagination_regressions import _node, _section


def question(number, space):
    return {'question': {'id': number, 'question_type': 'detailed_answer',
        'content': f'QUESTIONMARK{number} 已知函数，说明结论。', 'image_paths': []},
        'score': 10, 'solution_space': space}


@pytest.mark.parametrize('value,expected', [('30',30),('35',30),('-1',0),('17.5',17.5),
    (float('nan'),0),(float('inf'),0),('invalid',0)])
def test_space_limits_are_finite_and_explicit_zero_overrides_old_question_value(value, expected):
    assert normalize_solution_space_cm(value) == expected
    assert item_solution_space_cm({'solution_space':0}, {'solution_space':7}) == 0


def test_splittable_answer_space_uses_remaining_page_instead_of_expanding_question():
    _node("global.window={};\n" + _section('function paginatePaperBlocksByHeight(', 'function getPaperBlockOuterHeight(') + r'''
const assert=require('node:assert/strict');
const blocks=[{type:'section_title',qType:'detailed_answer',height:40},
 {type:'question',qType:'detailed_answer',questionId:1,height:1234,spaceHeight:1134,minimumFirstSpaceHeight:0},
 {type:'question',qType:'detailed_answer',questionId:2,height:100}];
const pages=paginatePaperBlocksByHeight(blocks,900,1000);
assert.equal(pages.length,2);
assert.deepEqual(pages.flat().filter(x=>x.type==='question').map(x=>x.questionId),[1,2]);
assert.ok(pages.every((p,i)=>p.reduce((h,b)=>h+b.height,0)<=(i?1000:900)));
const first=pages[0][1];assert.equal(first.allocatedSpaceHeight,760);
assert.equal(first.allocatedSpaceHeight+pages[1][0].height,1134);
// A real embedded figure stays intact at the top of the first answer space.
blocks[1].minimumFirstSpaceHeight=300;
const withFigure=paginatePaperBlocksByHeight(blocks,300,1000);
const owner=withFigure.flat().find(x=>x.questionId===1&&x.type==='question');
assert.ok(owner.allocatedSpaceHeight>=300);
assert.equal(withFigure.flat().filter(x=>x.type==='question').length,2);
''')


def test_word_long_space_is_exact_height_paragraphs_without_unbreakable_huge_spacing():
    data, _ = build_word_document('留白验证','','quiz',[question(1,'30'),question(2,'0')],
                                  show_secret=False,show_notice=False)
    doc = Document(BytesIO(data))
    blanks = [p for p in doc.paragraphs if not p.text
        and p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY]
    assert sum(p.paragraph_format.line_spacing.cm for p in blanks) == pytest.approx(30, abs=.02)
    assert all(p.paragraph_format.line_spacing.cm <= 2.01 for p in blanks)
    assert all(not p.paragraph_format.keep_with_next for p in blanks)
    assert all(not p.paragraph_format.space_after for p in blanks)
    assert 'QUESTIONMARK1' in '\n'.join(p.text for p in doc.paragraphs)
    assert 'QUESTIONMARK2' in '\n'.join(p.text for p in doc.paragraphs)


@pytest.mark.skipif(os.environ.get('MATHBANK_TEST_SPACING_NATIVE')!='1',reason='opt-in native XeLaTeX page validation')
def test_native_pdf_thirty_centimeters_continues_before_the_next_question(tmp_path):
    import pymupdf as fitz
    tex = build_latex_document('跨页留白验证','','exam',[question(1,'30'),question(2,'0')],
                                show_secret=False,show_notice=False)
    assert r'\vspace*{30.0cm}' not in tex
    (tmp_path/'space30.tex').write_text(tex)
    pdf, log = compile_tex_to_pdf(tex)
    (tmp_path/'space30.log').write_text(log)
    assert pdf, log
    (tmp_path/'space30.pdf').write_bytes(pdf)
    assert 'Overfull \\vbox' not in log
    doc = fitz.open(stream=pdf,filetype='pdf')
    assert len(doc)==2
    assert 'QUESTIONMARK1' in doc[0].get_text()
    assert 'QUESTIONMARK2' not in doc[0].get_text()
    assert 'QUESTIONMARK2' in doc[1].get_text()
    for page in doc:
        text=page.get_text()
        assert '数学试题第' in text.replace(' ','')
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                assert line['bbox'][3] < page.rect.height-15


@pytest.mark.skipif(os.environ.get('MATHBANK_TEST_SPACING_NATIVE')!='1',reason='opt-in native XeLaTeX image/page validation')
def test_native_embedded_figure_does_not_turn_thirty_cm_into_an_unbreakable_box(tmp_path):
    import pymupdf as fitz
    from PIL import Image, ImageDraw
    path=tmp_path/'figure.png'
    image=Image.new('RGB',(160,80),'white')
    ImageDraw.Draw(image).polygon([(10,70),(80,10),(150,70)],outline='black',width=3)
    image.save(path)
    first=question(1,'30')
    first['question'].update(content=first['question']['content']+'\n\n![](/static/uploads/figure.png)',
        image_paths=['/static/uploads/figure.png'],figure_align='bottom_right',figure_size='small')
    tex=build_latex_document('带图跨页留白','','quiz',[first,question(2,0)])
    assert '[30.0cm]' not in tex
    (tmp_path/'space30-figure.tex').write_text(tex)
    pdf,log=compile_tex_to_pdf(tex,[str(path)])
    assert pdf,log
    (tmp_path/'space30-figure.pdf').write_bytes(pdf)
    assert 'Overfull \\vbox' not in log
    doc=fitz.open(stream=pdf,filetype='pdf')
    assert len(doc)==2
    assert doc[0].get_images()
    assert 'QUESTIONMARK2' in doc[1].get_text()
